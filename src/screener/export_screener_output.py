"""
N100 Financial Intelligence Platform
Sprint 3, Day 17: Export screener_output.xlsx

One sheet per preset (6 total), 20 KPI columns, colour-coded:
green fill = cell meets that preset's threshold, red = fails it.
"""

import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__))))

from presets import PRESETS
import pandas as pd
from openpyxl.styles import PatternFill

OUTPUT_PATH = "output/screener_output.xlsx"

GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

DISPLAY_COLUMNS = [
    "company_id", "broad_sector", "return_on_equity_pct", "return_on_capital_employed_pct",
    "net_profit_margin_pct", "operating_profit_margin_pct", "debt_to_equity",
    "interest_coverage", "icr_label", "free_cash_flow_cr", "fcf_cagr_5yr",
    "revenue_cagr_5yr", "pat_cagr_5yr", "eps_cagr_5yr", "asset_turnover",
    "pe_ratio", "pb_ratio", "dividend_yield_pct", "dividend_payout_ratio_pct",
    "composite_quality_score",
]

# Preset-specific threshold checks, used to decide green vs red per cell.
# Maps preset name -> {column_name: (comparison_func, threshold)}
PRESET_THRESHOLDS = {
    "Quality Compounder": {
        "return_on_equity_pct": (lambda v, t: v > t, 15),
        "debt_to_equity": (lambda v, t: v < t, 1.0),
        "free_cash_flow_cr": (lambda v, t: v > t, 0),
        "revenue_cagr_5yr": (lambda v, t: v > t, 10),
    },
    "Value Pick": {
        "pe_ratio": (lambda v, t: v < t, 20),
        "pb_ratio": (lambda v, t: v < t, 3.0),
        "debt_to_equity": (lambda v, t: v < t, 2.0),
        "dividend_yield_pct": (lambda v, t: v > t, 1),
    },
    "Growth Accelerator": {
        "pat_cagr_5yr": (lambda v, t: v > t, 20),
        "revenue_cagr_5yr": (lambda v, t: v > t, 15),
        "debt_to_equity": (lambda v, t: v < t, 2.0),
    },
    "Dividend Champion": {
        "dividend_yield_pct": (lambda v, t: v > t, 2),
        "dividend_payout_ratio_pct": (lambda v, t: v < t, 80),
        "free_cash_flow_cr": (lambda v, t: v > t, 0),
    },
    "Debt-Free Blue Chip": {
        "debt_to_equity": (lambda v, t: v < t, 0.01),
        "return_on_equity_pct": (lambda v, t: v > t, 12),
    },
    "Turnaround Watch": {
        "revenue_cagr_5yr": (lambda v, t: True, None),  # revenue_cagr_3yr not in display cols; skip coloring
        "free_cash_flow_cr": (lambda v, t: v > t, 0),
    },
}


def write_sheet(writer, sheet_name, df):
    display_df = df[[c for c in DISPLAY_COLUMNS if c in df.columns]]
    display_df.to_excel(writer, sheet_name=sheet_name[:31], index=False)  # Excel sheet names cap at 31 chars

    worksheet = writer.sheets[sheet_name[:31]]
    thresholds = PRESET_THRESHOLDS.get(sheet_name, {})

    for col_idx, col_name in enumerate(display_df.columns, start=1):
        if col_name not in thresholds:
            continue
        compare_func, threshold = thresholds[col_name]
        if threshold is None:
            continue
        for row_idx, value in enumerate(display_df[col_name], start=2):  # row 1 is header
            if pd.isna(value):
                continue
            cell = worksheet.cell(row=row_idx, column=col_idx)
            try:
                if compare_func(value, threshold):
                    cell.fill = GREEN
                else:
                    cell.fill = RED
            except TypeError:
                continue


def write_notes_sheet(writer):
    """
    Explains the color-coding, and specifically the Financials-sector D/E
    exemption interaction: a company can legitimately appear on a preset
    sheet (via sector exemption) while its D/E cell shows red, because the
    cell coloring checks the raw numeric threshold with no sector awareness,
    while the filter itself correctly applies the exemption. Confirmed
    2026-08 via ICICIBANK and CANBK -- both included with red D/E cells.
    """
    note_df = pd.DataFrame({
        "Note": [
            "Green fill = cell value meets that preset's raw numeric threshold.",
            "Red fill = cell value does not meet the raw numeric threshold.",
            "",
            "IMPORTANT: Financials-sector companies may appear on a sheet with a red",
            "debt_to_equity cell -- this is correct, not a bug. Per spec, the D/E filter",
            "is exempted for Financials (high leverage is structurally normal for",
            "banks/NBFCs), so these companies qualify via sector exemption even though",
            "their raw D/E number fails the numeric threshold shown in red.",
        ]
    })
    note_df.to_excel(writer, sheet_name="Notes", index=False)


def main():
    os.makedirs("output", exist_ok=True)

    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        for preset_name, preset_func in PRESETS.items():
            print(f"Running {preset_name}...")
            result = preset_func()
            write_sheet(writer, preset_name, result)
            print(f"  {len(result)} companies written to sheet '{preset_name[:31]}'")

        write_notes_sheet(writer)
        print("  Notes sheet added")

    print(f"\n{OUTPUT_PATH} written with {len(PRESETS) + 1} sheets (6 presets + Notes).")


if __name__ == "__main__":
    main()